#!/usr/bin/env python3
"""Export the complete public résumé. Install scripts/resume-export-requirements.txt first."""
from pathlib import Path
import re
import shutil
from html import escape
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'Mostlylucid/Markdown/site/resume-complete.md'
DEST = ROOT / 'Mostlylucid/wwwroot/resumes'
DEST.mkdir(parents=True, exist_ok=True)
NAME = 'Scott-Galloway-Complete-Resume'


def plain(text):
    return re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text).replace('**', '')


def pdf_markup(text):
    text = escape(text)
    return re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'<a href="\2" color="#087c7a">\1</a>', text)


doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.75)
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'Calibri', Pt(10)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12
for name, size in [('Title', 28), ('Heading 1', 16), ('Heading 2', 12)]:
    style = doc.styles[name]
    style.font.name, style.font.size = 'Calibri', Pt(size)
    style.font.color.rgb = RGBColor.from_string('103C45')
    style.paragraph_format.keep_with_next = True
footer = section.footer.paragraphs[0]
footer.text = 'Scott Galloway · Complete résumé · mostlylucid.net'
footer.runs[0].font.size = Pt(8)
footer.runs[0].font.color.rgb = RGBColor.from_string('526674')
doc.core_properties.title = 'Scott Galloway — Complete Résumé'
doc.core_properties.author = 'Scott Galloway'

styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name='ResumeTitle', fontName='Helvetica-Bold', fontSize=26, leading=31, textColor=colors.HexColor('#103c45'), spaceAfter=14, keepWithNext=True))
styles.add(ParagraphStyle(name='ResumeH2', fontName='Helvetica-Bold', fontSize=15, leading=19, textColor=colors.HexColor('#103c45'), spaceBefore=16, spaceAfter=9, keepWithNext=True))
styles.add(ParagraphStyle(name='ResumeH3', fontName='Helvetica-Bold', fontSize=11, leading=15, textColor=colors.HexColor('#112f43'), spaceBefore=12, spaceAfter=7, keepWithNext=True))
styles.add(ParagraphStyle(name='ResumeBody', fontName='Helvetica', fontSize=9.5, leading=14, spaceAfter=8, alignment=TA_LEFT))
styles.add(ParagraphStyle(name='ResumeBullet', parent=styles['ResumeBody'], leftIndent=10, firstLineIndent=-8, spaceAfter=6))
flow = []
for line in SOURCE.read_text().splitlines():
    if not line.strip():
        continue
    if line.startswith('# '):
        text = line[2:]
        doc.add_heading(plain(text), 0)
        flow.append(Paragraph(pdf_markup(text), styles['ResumeTitle']))
    elif line.startswith('## '):
        text = line[3:]
        doc.add_heading(plain(text), 1)
        flow.append(Paragraph(pdf_markup(text), styles['ResumeH2']))
    elif line.startswith('### '):
        text = line[4:]
        doc.add_heading(plain(text), 2)
        flow.append(Paragraph(pdf_markup(text), styles['ResumeH3']))
    elif line.startswith('- '):
        text = line[2:]
        doc.add_paragraph(plain(text), style='List Bullet')
        flow.append(Paragraph('• ' + pdf_markup(text), styles['ResumeBullet']))
    else:
        doc.add_paragraph(plain(line))
        flow.append(Paragraph(pdf_markup(line), styles['ResumeBody']))


def footer_pdf(canvas, document):
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor('#dce5e8'))
    canvas.line(19*mm, 16*mm, A4[0]-19*mm, 16*mm)
    canvas.setFont('Helvetica', 8)
    canvas.setFillColor(colors.HexColor('#526674'))
    canvas.drawString(19*mm, 11*mm, 'Scott Galloway · Complete résumé · mostlylucid.net')
    canvas.drawRightString(A4[0]-19*mm, 11*mm, str(document.page))
    canvas.restoreState()


doc.save(DEST / (NAME + '.docx'))
pdf = SimpleDocTemplate(str(DEST / (NAME + '.pdf')), pagesize=A4,
    leftMargin=19*mm, rightMargin=19*mm, topMargin=18*mm, bottomMargin=22*mm,
    title='Scott Galloway — Complete Résumé', author='Scott Galloway')
pdf.build(flow, onFirstPage=footer_pdf, onLaterPages=footer_pdf)
shutil.copyfile(SOURCE, DEST / (NAME + '.md'))
print(f'Exported PDF, DOCX and Markdown to {DEST}')

# Keep the bundled compiler source aligned with the complete public résumé.
from build_resume_record import write_record
write_record()
